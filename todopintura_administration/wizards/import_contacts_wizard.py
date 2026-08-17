import base64
import io
import logging

from markupsafe import escape

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

_logger = logging.getLogger(__name__)

PAYMENT_TERMS = {
    '1010': 'GIRO A 30 DIAS',
    '1011': 'GIRO 45 DIAS F.F.',
    '1012': 'GIRO 60 DIAS F.F.',
    '1013': 'GIRO 90 DIAS F.F.',
    '1014': 'GIRO 120 DIAS F.F.',
    '1020': 'GIRO 30-60 DIAS F.F.',
    '1030': 'GIRO 30-60-90 DIAS F.F.',
    '1031': 'GIRO 60-90-120 DIAS F.F.',
    '1040': 'GIRO 30-60-90-120 DIAS F.F.',
    '1041': 'GIRO 60-90-120-150 DIAS F.F.',
    '1021': 'GIRO 60-90 DIAS F.F.',
    '3000': 'CONTADO',
    '4010': 'REPOSICION PAGARE VTO. 30 DIAS FECHA FCT',
    '4020': 'REPOSICION PAGARES VTOS. 30/60 DIAS F.F.',
    '1015': 'GIRO 40 DIAS F.F.',
    '1022': 'GIRO 15-30 DIAS F.F.',
    '2010': 'GIRO PTE. ACEPT. VTO. 30 DIAS F.F.',
    '2011': 'GIRO PTE. ACEPT. VTO. 45 DIAS F.F.',
    '2012': 'GIRO PTE. ACEPT. VTO. 60 DIAS F.F.',
    '2013': 'GIRO PTE. ACEPT. VTO. 90 DIAS F.F.',
    '2014': 'GIRO PTE. ACEPT. VTO. 120 DIAS F.F.',
    '2015': 'GIRO PTE. ACEPT. VTO. 40 DIAS F.F.',
    '2020': 'GIRO PTE. ACEPT. VTO. 30-60 DIAS F.F.',
    '2022': 'GIRO PTE. ACEPT. VTO. 15-30 DIAS F.F.',
    '2030': 'GIRO PTE. ACEPT. VTO. 30-60-90 DIAS F.F.',
    '2031': 'GIRO PTE. ACEPT. VTO. 60-90-120 DIAS F.F',
    '2040': 'GIRO PTE. ACEPT. VTO. 30-60-90-120 DIAS',
    '2041': 'GIRO PTE. ACEPT. VTO. 60-90-120-150 DIAS',
    '2021': 'GIRO PTE. ACEPT. VTO. 60-90 DIAS F.F.',
    '4000': 'REPOSICION',
    '3020': 'CONTRA COMPROMISO',
    '4030': 'REPOSICION PAGARES VTOS. 30/60/90 D. F.F',
    '1016': 'GIRO 10 DIAS F.F.',
    '2032': 'GIRO PTE. ACEPT. VTO. 120-150-180 DIAS',
    '2023': 'GIRO PTE. ACEPT. VTO. 90-120 DIAS F.F.',
    '2033': 'GIRO PTE. ACEPT. VTO. 90-120-150 DIAS',
    '2016': 'GIRO PTE. ACEPT. VTO. 75 DIAS F.F.',
    '2052': 'GIRO PTE. ACEPT. 60-90-120-150-180 DIAS',
    '1033': 'GIRO 90-120-150 DIAS F.F.',
    '2017': 'GIRO PTE. ACEPT. VTO. 180 DIAS F.F.',
    '4040': 'REPOSIC. PAGARES VTOS. 30/60/90/120 F.F.',
    '4011': 'REPOSICION PAGARE VTO. 45 DIAS FECHA FCT',
    '4012': 'REPOSICION PAGARE VTO. 60 DIAS FECHA FCT',
    '4013': 'REPOSICION PAGARE VTO. 85 DIAS FECHA FCT',
    '4014': 'REPOSICION PAGARE VTO. 120 DIAS FECHA F.',
    '4021': 'REPOSICION PAGARES VTOS. 60/90 DIAS F.F.',
    '4022': 'REPOSICION PAGARES VTOS. 15/30 DIAS F.F.',
    '4031': 'REPOSICION PAGARES VTOS. 60/90/120  F.F.',
    '4050': 'TRANSFERENCIA   ES5721009753822200083736',
    '4999': 'CLIENTE OBSOLETO',
    '1017': 'GIRO 150 DIAS F.F.',
    '1023': 'GIRO 90-120 DIAS F.F.',
    '3011': 'RECIBO AL COBRO 15 D. FECHA FCT.VENDEDOR',
    '3012': 'RECIBO AL COBRO 30 D. FECHA FCT.VENDEDOR',
    '4510': 'CONFIRMING 240 DIAS',
    '4501': 'CONFIRMING 60 DIAS',
    '4502': 'CONFIRMING 90 DIAS',
    '4503': 'CONFIRMING 120 DIAS',
    '4504': 'CONFIRMING 180 DIAS',
    '3018': 'REPOSICION TALON BANCARIO 60 DIAS F.F.',
    '3019': 'REPOSICION TALON BANCARIO 90 DIAS F.F.',
    '5000': 'ABONOS AUTOMATICOS',
    '5999': 'FACTURA RECTIFICATIVA ABONO',
    '4505': 'CONFIRMING 150 DIAS',
    '4055': 'TRANSF.365 DIAS ES5721009753822200083736',
    '4015': 'REPOSICION PAGARES VTO. 150 DIAS F.F',
    '1026': 'GIRO 45-90 DIAS F.F.',
    '4041': 'REPOSICION PAGARES VTO: 90/120 F.F.',
    '4506': 'CONFIRMING 45 DIAS',
    '4018': 'REPOSICION PAGARES VTOS.75 DIAS F.F.',
    '1019': 'GIRO 75 DIAS F.F.',
    '3021': 'REPOSICION TALON BANCARIO 75 DIAS F.F',
    '4507': 'CONFIRMING 85 DIAS',
    '1034': 'GIRO 30-60-75 DIAS F.F.',
    '4056': 'TRANSF.15 DIAS ES5721009753822200083736',
    '4019': 'REPOSICION PAGARE VTO. 90 DIAS FECHA FCT',
    '3010': '',
    '3040': '',
    '4057': 'TRANSF.75 DIAS ES5721009753822200083736',
    '4508': 'CONFIRMING 30 60 90',
    '4058': 'TRANSF.85 DIAS ES5721009753822200083736',
    '3500': 'COMPROMISO COMPRAS',
    '4059': 'TRANSF/45 DIAS  ES5721009753822200083736',
}


class ImportContactsWizard(models.TransientModel):
    _name = 'import.contacts.wizard'
    _description = 'Wizard para importar contactos desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX')
    file_name = fields.Char('Nombre del archivo')
    state = fields.Selection([
        ('upload', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='upload', required=True)
    log = fields.Html('Registro de importación', readonly=True)
    created_count = fields.Integer('Clientes nuevos', readonly=True)
    updated_count = fields.Integer('Clientes actualizados', readonly=True)
    error_count = fields.Integer('Errores', readonly=True)

    @api.model
    def _sanitize_import_text(self, value):
        if value is None:
            return ''
        text = value.strip() if isinstance(value, str) else str(value).strip()
        return '' if text and all(char == '*' for char in text) else text

    @api.model
    def _normalize_contact_vat(self, nif):
        nif = self._sanitize_import_text(nif).upper()
        if not nif:
            return '', 'ES'

        nif = ''.join(char for char in nif if char.isalnum())
        country_code = nif[:2] if nif[:2].isalpha() else 'ES'
        vat_number = nif[2:] if nif[:2].isalpha() else nif
        if not vat_number:
            return '', country_code
        return f"{country_code}{vat_number}", country_code

    @api.model
    def _normalize_iban(self, iban):
        iban = self._sanitize_import_text(iban).upper()
        normalized_iban = ''.join(char for char in iban if char.isalnum())
        if not normalized_iban:
            return ''

        if normalized_iban in {'NONE', 'NULL', 'FALSE', 'NA', 'NAN'}:
            return ''

        if set(normalized_iban) == {'0'}:
            return ''

        return normalized_iban

    @api.model
    def _parse_credit_limit(self, value):
        if value in (None, 0, 0.0, False):
            return None
        if isinstance(value, (int, float)):
            return value

        value = self._sanitize_import_text(value)
        if not value:
            return None

        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    @api.model
    def _find_existing_contact(self, num_client, name, vat, iban=None):
        partner_model = self.env['res.partner'].with_context(active_test=False)
        base_domain = [('parent_id', '=', False), ('is_company', '=', True)]

        search_domains = []
        if num_client not in ['', None]:
            search_domains.append(base_domain + [('ref', '=', num_client)])
        if vat:
            search_domains.append(base_domain + [('vat', '=', vat)])
        if num_client not in ['', None] and name:
            search_domains.append(base_domain + [('ref', '=', num_client), ('name', '=', name)])
        if vat and name:
            search_domains.append(base_domain + [('vat', '=', vat), ('name', '=', name)])

        for domain in search_domains:
            contact = partner_model.search(domain, limit=1)
            if contact:
                return contact

        normalized_iban = self._normalize_iban(iban)
        if normalized_iban:
            bank_record = self.env['res.partner.bank'].with_context(active_test=False).search([
                ('acc_number', '=', normalized_iban),
            ], limit=1)
            if bank_record:
                return bank_record.partner_id.commercial_partner_id

        return partner_model.browse()

    @api.model
    def _create_or_update_contact(self, num_client, name, record, iban=None):
        """Crea el contacto si no existe, o actualiza sus datos si ya existe.

        Devuelve una tupla (contact, status) donde status es 'created' o 'updated',
        para poder reflejarlo en el log de importación.
        """
        contact = self._find_existing_contact(num_client, name, record.get('vat'), iban=iban)

        try:
            if contact:
                contact.write(record)
                status = 'updated'
            else:
                contact = self.env['res.partner'].create(record)
                status = 'created'
        except Exception:
            _logger.exception("Error al crear/actualizar el contacto %s, reintentando sin VAT", name)
            fallback_record = dict(record, vat='')
            contact = contact or self._find_existing_contact(num_client, name, record.get('vat'), iban=iban)
            if contact:
                contact.write(fallback_record)
                status = 'updated'
            else:
                contact = self.env['res.partner'].create(fallback_record)
                status = 'created'

        return contact, status

    @api.model
    def _ensure_bank_account(self, contact, iban):
        normalized_iban = self._normalize_iban(iban)
        if not normalized_iban or not contact:
            return

        bank_model = self.env['res.partner.bank'].with_context(active_test=False)
        existing_bank_record = bank_model.search([
            ('acc_number', '=', normalized_iban),
            ('partner_id', '=', contact.id)
        ], limit=1)

        if not existing_bank_record:
            existing_bank_record = bank_model.search([
                ('acc_number', '=', normalized_iban),
            ], limit=1)

        if not existing_bank_record:
            self.env['res.partner.bank'].create({
                'acc_number': normalized_iban,
                'partner_id': contact.id
            })

    @api.model
    def _prepare_contact_record(self, num_client, name, address, cp, telefono, vat, email, notes, country_id):
        return {
            'ref': num_client,
            'name': name,
            'street': address,
            'zip': cp,
            'country_id': country_id.id,
            'phone': telefono,
            'vat': vat,
            'email': email,
            'comment': notes,
            'is_company': True,
        }

    @api.model
    def _iter_rows(self, sheet, is_xlsx):
        if is_xlsx:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                yield list(row)
        else:
            for row_index in range(1, sheet.nrows):
                yield [sheet.cell(row_index, col).value for col in range(sheet.ncols)]

    @api.model
    def _process_import_row(self, row, country_id, results):
        """Procesa una fila del Excel: crea/actualiza el cliente y registra el resultado.

        Devuelve False cuando la fila está completamente vacía (fin de los datos reales),
        True en caso contrario.
        """
        def cell(idx):
            return row[idx] if len(row) > idx else None

        num_client = int(cell(0)) if cell(0) else ''
        name = self._sanitize_import_text(cell(1))
        address = self._sanitize_import_text(cell(2))
        cp_value = cell(3)
        cp = self._sanitize_import_text(str(int(cp_value)) if cp_value and isinstance(cp_value, (int, float)) else '')
        telefono_value = cell(4)
        telefono = self._sanitize_import_text(
            str(int(telefono_value)) if telefono_value and isinstance(telefono_value, (int, float)) else '')
        nif = self._sanitize_import_text(cell(6))
        try:
            forma_pago = str(int(cell(8))) if cell(8) else ''
        except (ValueError, TypeError):
            forma_pago = ''
        email = self._sanitize_import_text(cell(23))
        credit_limit = self._parse_credit_limit(cell(11))
        iban = self._sanitize_import_text(cell(14))
        observation1 = self._sanitize_import_text(cell(19))
        observation2 = self._sanitize_import_text(cell(20))
        observation3 = self._sanitize_import_text(cell(21))
        observation4 = self._sanitize_import_text(cell(22))

        if not (name or address or cp or telefono or nif):
            return False

        row_label = f"[{num_client or 's/n'}] {name or '(sin nombre)'}"

        try:
            notes = "<br/>".join(filter(None, [observation1, observation2, observation3, observation4]))
            vat, _country_code = self._normalize_contact_vat(nif)
            record = self._prepare_contact_record(num_client, name, address, cp, telefono, vat, email, notes, country_id)

            if credit_limit is not None:
                record['use_partner_credit_limit'] = True
                record['credit_limit'] = credit_limit
            if forma_pago in PAYMENT_TERMS and PAYMENT_TERMS[forma_pago]:
                payment_term = self.env['account.payment.term'].search(
                    [('name', '=', PAYMENT_TERMS[forma_pago])], limit=1)
                if payment_term:
                    record['property_payment_term_id'] = payment_term.id

            contact, status = self._create_or_update_contact(num_client, name, record, iban=iban)
            self.env.cr.flush()
            self._ensure_bank_account(contact, iban)
            results[status].append(row_label)
        except Exception as exc:
            _logger.exception("Error al importar el cliente %s", row_label)
            results['error'].append(f"{row_label}: {exc}")

        return True

    @api.model
    def _build_import_log_html(self, results):
        def section(title, items, css_class):
            if not items:
                return f'<p><b>{escape(title)}:</b> ninguno.</p>'
            lines = ''.join(f'<li>{escape(line)}</li>' for line in items)
            return f'<p><b>{escape(title)} ({len(items)}):</b></p><ul class="{css_class}">{lines}</ul>'

        return (
            section('Clientes nuevos', results['created'], 'text-success')
            + section('Clientes actualizados', results['updated'], 'text-info')
            + section('Errores', results['error'], 'text-danger')
        )

    def _load_sheet(self):
        ext = ''
        if self.file_name:
            ext = self.file_name.split('.')[-1].lower()
        data = base64.b64decode(self.file)
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'

        if ext == 'xlsx':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return wb.active, True
        elif ext == 'xls':
            if not xlrd:
                raise UserError("Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.")
            book = xlrd.open_workbook(file_contents=data)
            return book.sheet_by_index(0), False
        else:
            raise UserError("Formato de archivo no soportado. Usa .xls o .xlsx")

    def action_import_contacts(self):
        self.ensure_one()
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        sheet, is_xlsx = self._load_sheet()

        country_id = self.env['res.country'].search([('code', '=', 'ES')], limit=1)
        if not country_id:
            raise UserError("País 'España' no encontrado en la base de datos.")

        results = {'created': [], 'updated': [], 'error': []}
        for row in self._iter_rows(sheet, is_xlsx):
            if not self._process_import_row(row, country_id, results):
                break

        self.write({
            'state': 'done',
            'created_count': len(results['created']),
            'updated_count': len(results['updated']),
            'error_count': len(results['error']),
            'log': self._build_import_log_html(results),
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.contacts.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_reset(self):
        self.ensure_one()
        self.write({
            'state': 'upload',
            'file': False,
            'file_name': False,
            'log': False,
            'created_count': 0,
            'updated_count': 0,
            'error_count': 0,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.contacts.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
