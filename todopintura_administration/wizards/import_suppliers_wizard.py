from odoo import api, fields, models
import base64
from odoo.exceptions import UserError
import io
try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None


class ImportSuppliersWizard(models.TransientModel):
    _name = 'import.suppliers.wizard'
    _description = 'Wizard para importar proveedores desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX', required=True)
    file_name = fields.Char('Nombre del archivo')

    @api.model
    def _sanitize_import_text(self, value):
        if value is None:
            return ''
        text = value.strip() if isinstance(value, str) else str(value).strip()
        return '' if text and all(char == '*' for char in text) else text

    @api.model
    def _normalize_supplier_vat(self, nif):
        nif = self._sanitize_import_text(nif).upper()
        if not nif:
            return '', 'ES'

        nif = ''.join(char for char in nif if char.isalnum())
        country_code = nif[:2] if nif[:2] in {'ES', 'DE', 'GB'} else 'ES'
        vat_number = nif[2:] if nif[:2] in {'ES', 'DE', 'GB'} else nif
        return f"{country_code}{vat_number}" if vat_number else '', country_code

    @api.model
    def _find_existing_supplier(self, num_prov, name, vat):
        partner_model = self.env['res.partner'].with_context(active_test=False)
        base_domain = [('parent_id', '=', False), ('is_company', '=', True)]

        search_domains = []
        if num_prov:
            search_domains.append(base_domain + [('ref', '=', num_prov)])
        if vat:
            search_domains.append(base_domain + [('vat', '=', vat)])
        if num_prov and name:
            search_domains.append(base_domain + [('ref', '=', num_prov), ('name', '=', name)])
        if vat and name:
            search_domains.append(base_domain + [('vat', '=', vat), ('name', '=', name)])

        for domain in search_domains:
            supplier = partner_model.search(domain, limit=1)
            if supplier:
                return supplier

        return partner_model.browse()

    @api.model
    def _create_or_update_supplier(self, num_prov, name, record):
        supplier = self._find_existing_supplier(num_prov, name, record.get('vat'))

        try:
            if supplier:
                supplier.write(record)
                print(f"Proveedor actualizado: {supplier.name}")
            else:
                supplier = self.env['res.partner'].create(record)
                print(f"Proveedor creado: {name}")
        except Exception as e:
            print(f"Error al actualizar o crear el proveedor: {e}")
            fallback_record = dict(record, vat='')
            supplier = supplier or self._find_existing_supplier(num_prov, name, record.get('vat'))
            if supplier:
                supplier.write(fallback_record)
            else:
                supplier = self.env['res.partner'].create(fallback_record)

        return supplier

    @api.model
    def _upsert_secondary_address(self, supplier, name, address2, cp2):
        if not (address2 or cp2):
            return

        contact_address = {
            'name': "Otra dirección " + str(name),
            'parent_id': supplier.id,
            'type': 'other',
            'street': address2,
            'zip': cp2,
        }
        existing_contact = self.env['res.partner'].with_context(active_test=False).search([
            ('name', '=', "Otra dirección " + str(name)),
            ('parent_id', '=', supplier.id),
            ('type', '=', 'other'),
            ('street', '=', address2),
            ('zip', '=', cp2)
        ], limit=1)

        if existing_contact:
            existing_contact.write(contact_address)
            print(
                f"Dirección secundaria actualizada: {address2}, Nombre: {existing_contact.name}, Parent ID: {existing_contact.parent_id.id}")
        else:
            self.env['res.partner'].create(contact_address)
            print(
                f"Dirección secundaria creada: {address2}, Nombre: {contact_address['name']}, Parent ID: {contact_address['parent_id']}")

    def action_import_suppliers(self):
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        payment_terms = {
            '1010': 'GIRO A 30 DIAS',
            '1011': 'GIRO 45 DIAS F.F.',
            '1012': 'GIRO 60 DIAS F.F.',
            '1013': 'GIRO 90 DIAS F.F.',
            '1014': 'GIRO 120 DIAS F.F.',
            '1020': 'GIRO 30-60 DIAS F.F.',
            '1030': 'GIRO 30-60-90 DIAS F.F.',
            '1031': 'GIRO 60-90-120 DIAS F.F.',
            '1040': 'GIRO 30-60-90-120 DIAS F.F.',
            '1041': 'GIRO 60-90-120-150 DIAS F.F.',
            '1021': 'GIRO 60-90 DIAS F.F.',
            '3000': 'CONTADO',
            '4010': 'REPOSICION PAGARE VTO. 30 DIAS FECHA FCT',
            '4020': 'REPOSICION PAGARES VTOS. 30/60 DIAS F.F.',
            '1015': 'GIRO 40 DIAS F.F.',
            '1022': 'GIRO 15-30 DIAS F.F.',
            '2010': 'GIRO PTE. ACEPT. VTO. 30 DIAS F.F.',
            '2011': 'GIRO PTE. ACEPT. VTO. 45 DIAS F.F.',
            '2012': 'GIRO PTE. ACEPT. VTO. 60 DIAS F.F.',
            '2013': 'GIRO PTE. ACEPT. VTO. 90 DIAS F.F.',
            '2014': 'GIRO PTE. ACEPT. VTO. 120 DIAS F.F.',
            '2015': 'GIRO PTE. ACEPT. VTO. 40 DIAS F.F.',
            '2020': 'GIRO PTE. ACEPT. VTO. 30-60 DIAS F.F.',
            '2022': 'GIRO PTE. ACEPT. VTO. 15-30 DIAS F.F.',
            '2030': 'GIRO PTE. ACEPT. VTO. 30-60-90 DIAS F.F.',
            '2031': 'GIRO PTE. ACEPT. VTO. 60-90-120 DIAS F.F',
            '2040': 'GIRO PTE. ACEPT. VTO. 30-60-90-120 DIAS',
            '2041': 'GIRO PTE. ACEPT. VTO. 60-90-120-150 DIAS',
            '2021': 'GIRO PTE. ACEPT. VTO. 60-90 DIAS F.F.',
            '4000': 'REPOSICION',
            '3020': 'CONTRA COMPROMISO',
            '4030': 'REPOSICION PAGARES VTOS. 30/60/90 D. F.F',
            '1016': 'GIRO 10 DIAS F.F.',
            '2032': 'GIRO PTE. ACEPT. VTO. 120-150-180 DIAS',
            '2023': 'GIRO PTE. ACEPT. VTO. 90-120 DIAS F.F.',
            '2033': 'GIRO PTE. ACEPT. VTO. 90-120-150 DIAS',
            '2016': 'GIRO PTE. ACEPT. VTO. 75 DIAS F.F.',
            '2052': 'GIRO PTE. ACEPT. 60-90-120-150-180 DIAS',
            '1033': 'GIRO 90-120-150 DIAS F.F.',
            '2017': 'GIRO PTE. ACEPT. VTO. 180 DIAS F.F.',
            '4040': 'REPOSIC. PAGARES VTOS. 30/60/90/120 F.F.',
            '4011': 'REPOSICION PAGARE VTO. 45 DIAS FECHA FCT',
            '4012': 'REPOSICION PAGARE VTO. 60 DIAS FECHA FCT',
            '4013': 'REPOSICION PAGARE VTO. 85 DIAS FECHA FCT',
            '4014': 'REPOSICION PAGARE VTO. 120 DIAS FECHA F.',
            '4021': 'REPOSICION PAGARES VTOS. 60/90 DIAS F.F.',
            '4022': 'REPOSICION PAGARES VTOS. 15/30 DIAS F.F.',
            '4031': 'REPOSICION PAGARES VTOS. 60/90/120  F.F.',
            '4050': 'TRANSFERENCIA   ES5721009753822200083736',
            '4999': 'CLIENTE OBSOLETO',
            '1017': 'GIRO 150 DIAS F.F.',
            '1023': 'GIRO 90-120 DIAS F.F.',
            '3011': 'RECIBO AL COBRO 15 D. FECHA FCT.VENDEDOR',
            '3012': 'RECIBO AL COBRO 30 D. FECHA FCT.VENDEDOR',
            '4510': 'CONFIRMING 240 DIAS',
            '4501': 'CONFIRMING 60 DIAS',
            '4502': 'CONFIRMING 90 DIAS',
            '4503': 'CONFIRMING 120 DIAS',
            '4504': 'CONFIRMING 180 DIAS',
            '3018': 'REPOSICION TALON BANCARIO 60 DIAS F.F.',
            '3019': 'REPOSICION TALON BANCARIO 90 DIAS F.F.',
            '5000': 'ABONOS AUTOMATICOS',
            '5999': 'FACTURA RECTIFICATIVA ABONO',
            '4505': 'CONFIRMING 150 DIAS',
            '4055': 'TRANSF.365 DIAS ES5721009753822200083736',
            '4015': 'REPOSICION PAGARES VTO. 150 DIAS F.F',
            '1026': 'GIRO 45-90 DIAS F.F.',
            '4041': 'REPOSICION PAGARES VTO: 90/120 F.F.',
            '4506': 'CONFIRMING 45 DIAS',
            '4018': 'REPOSICION PAGARES VTOS.75 DIAS F.F.',
            '1019': 'GIRO 75 DIAS F.F.',
            '3021': 'REPOSICION TALON BANCARIO 75 DIAS F.F',
            '4507': 'CONFIRMING 85 DIAS',
            '1034': 'GIRO 30-60-75 DIAS F.F.',
            '4056': 'TRANSF.15 DIAS ES5721009753822200083736',
            '4019': 'REPOSICION PAGARE VTO. 90 DIAS FECHA FCT',
            '3010': '',
            '3040': '',
            '4057': 'TRANSF.75 DIAS ES5721009753822200083736',
            '4508': 'CONFIRMING 30 60 90',
            '4058': 'TRANSF.85 DIAS ES5721009753822200083736',
            '3500': 'COMPROMISO COMPRAS',
            '4059': 'TRANSF/45 DIAS  ES5721009753822200083736'
        }

        ext = ''
        if self.file_name:
            ext = self.file_name.split('.')[-1].lower()
        data = base64.b64decode(self.file)
        # Detección por cabecera si la extensión no es fiable
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'
        sheet = None
        is_xlsx = False
        if ext == 'xlsx':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.")
            is_xlsx = True
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = wb.active
        elif ext == 'xls':
            if not xlrd:
                raise UserError("Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.")
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
        else:
            raise UserError("Formato de archivo no soportado. Usa .xls o .xlsx")

        if is_xlsx:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                num_prov = '0' + str(int(row[0])) if row[0] else ''
                name = self._sanitize_import_text(row[1])
                address = self._sanitize_import_text(row[2])
                telefono_value = row[3]
                telefono = str(int(telefono_value)) if telefono_value and isinstance(telefono_value, (int, float)) else ''
                telefono = self._sanitize_import_text(telefono)
                telefono2_value = row[4]
                telefono2 = str(int(telefono2_value)) if telefono2_value and isinstance(telefono2_value, (int, float)) else ''
                telefono2 = self._sanitize_import_text(telefono2)
                nif_value = row[5]
                nif = self._sanitize_import_text(nif_value)
                forma_pago = str(int(row[8])) if row[8] else ''
                cp_value = row[7]
                cp = str(int(cp_value)) if cp_value and isinstance(cp_value, (int, float)) else ''
                cp = self._sanitize_import_text(cp)
                address2 = self._sanitize_import_text(row[9])
                cp2_value = row[10]
                cp2 = str(int(cp2_value)) if cp2_value and isinstance(cp2_value, (int, float)) else ''
                cp2 = self._sanitize_import_text(cp2)
                activo = row[15] if len(row) > 15 else ''
                # Observaciones
                observations = [str(row[i]).strip() if len(row) > i and row[i] is not None else '' for i in range(16, 40)]
                notes = "<br/>".join(filter(None, observations))

                if not (name or address or cp or telefono or nif):
                    print("Todos los datos están vacíos. Terminando la importación.")
                    break

                vat, country_code = self._normalize_supplier_vat(nif)
                es_country = self.env['res.country'].search([('code', '=', country_code)], limit=1)
                print(es_country.id, es_country.name)

                record = {
                    'ref': num_prov,
                    'name': name,
                    'street': address,
                    'zip': cp,
                    'is_company': True,
                    'country_id': es_country.id,
                    'phone': telefono,
                    'vat': vat,
                    'active': False if activo == 'N' else True,
                    'comment': notes,
                }

                if forma_pago in payment_terms:
                    print(f"Forma de pago encontrada: {forma_pago} - {payment_terms[forma_pago]}")
                    payment_term = self.env['account.payment.term'].search(
                        [('name', '=', payment_terms[forma_pago])], limit=1)
                    if payment_term:
                        print(f"Término de pago encontrado: {payment_term.name} (ID: {payment_term.id})")
                        record['property_supplier_payment_term_id'] = payment_term.id
                    else:
                        print(f"No se encontró un término de pago para: {payment_terms[forma_pago]}")

                supplier = self._create_or_update_supplier(num_prov, name, record)
                self._upsert_secondary_address(supplier, name, address2, cp2)
        else:
            # ...existing code para xlrd (xls)...
            data = base64.b64decode(self.file)
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
            for row in range(1, sheet.nrows):
                num_prov = '0' + str(int(sheet.cell(row, 0).value))
                name = self._sanitize_import_text(sheet.cell(row, 1).value)
                address = self._sanitize_import_text(sheet.cell(row, 2).value)
                telefono_value = sheet.cell(row, 3).value
                telefono = str(int(telefono_value)) if telefono_value and isinstance(telefono_value, (int, float)) else ''
                telefono = self._sanitize_import_text(telefono)
                telefono2_value = sheet.cell(row, 4).value
                telefono2 = str(int(telefono2_value)) if telefono2_value and isinstance(telefono2_value, (int, float)) else ''
                telefono2 = self._sanitize_import_text(telefono2)
                nif_value = sheet.cell(row, 5).value
                nif = self._sanitize_import_text(nif_value)
                forma_pago = str(int(sheet.cell(row, 8).value)) if sheet.cell(row, 8).value else ''
                cp_value = sheet.cell(row, 7).value
                cp = str(int(cp_value)) if cp_value and isinstance(cp_value, (int, float)) else ''
                cp = self._sanitize_import_text(cp)
                address2 = self._sanitize_import_text(sheet.cell(row, 9).value)
                cp2_value = sheet.cell(row, 10).value
                cp2 = str(int(cp2_value)) if cp2_value and isinstance(cp2_value, (int, float)) else ''
                cp2 = self._sanitize_import_text(cp2)
                activo = sheet.cell(row, 15).value
                observations = [str(sheet.cell(row, i).value).strip() if sheet.cell(row, i).value is not None else '' for i in range(16, 40)]
                notes = "<br/>".join(filter(None, observations))

                if not (name or address or cp or telefono or nif):
                    print("Todos los datos están vacíos. Terminando la importación.")
                    break

                vat, country_code = self._normalize_supplier_vat(nif)
                es_country = self.env['res.country'].search([('code', '=', country_code)], limit=1)
                print(es_country.id, es_country.name)

                record = {
                    'ref': num_prov,
                    'name': name,
                    'street': address,
                    'zip': cp,
                    'is_company': True,
                    'country_id': es_country.id,
                    'phone': telefono,
                    'vat': vat,
                    'active': False if activo == 'N' else True,
                    'comment': notes,
                }

                if forma_pago in payment_terms:
                    print(f"Forma de pago encontrada: {forma_pago} - {payment_terms[forma_pago]}")
                    payment_term = self.env['account.payment.term'].search(
                        [('name', '=', payment_terms[forma_pago])], limit=1)
                    if payment_term:
                        print(f"Término de pago encontrado: {payment_term.name} (ID: {payment_term.id})")
                        record['property_supplier_payment_term_id'] = payment_term.id
                    else:
                        print(f"No se encontró un término de pago para: {payment_terms[forma_pago]}")

                supplier = self._create_or_update_supplier(num_prov, name, record)
                self._upsert_secondary_address(supplier, name, address2, cp2)
