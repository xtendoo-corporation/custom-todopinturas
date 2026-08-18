import base64
import io
import logging

from markupsafe import escape

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

_logger = logging.getLogger(__name__)


class ImportPaymentTermsWizard(models.TransientModel):
    _name = 'import.payment.terms.wizard'
    _description = 'Wizard para importar formas de pago desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX')
    file_name = fields.Char('Nombre del archivo')
    state = fields.Selection([
        ('upload', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='upload', required=True)
    log = fields.Html('Registro de importación', readonly=True)
    created_count = fields.Integer('Formas de pago nuevas', readonly=True)
    updated_count = fields.Integer('Formas de pago actualizadas', readonly=True)
    error_count = fields.Integer('Errores', readonly=True)

    @api.model
    def _sanitize_import_text(self, value):
        if value is None:
            return ''
        text = value.strip() if isinstance(value, str) else str(value).strip()
        return '' if text and all(char == '*' for char in text) else text

    @api.model
    def _safe_int(self, value):
        if value in (None, False, ''):
            return None
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None

    @api.model
    def _find_existing_payment_term(self, codigo):
        return self.env['account.payment.term'].with_context(active_test=False).search(
            [('referencia_todopintura', '=', codigo)], limit=1)

    @api.model
    def _prepare_payment_term_record(self, codigo, name, condiciones, metodo_pago):
        note_lines = []
        if condiciones:
            note_lines.append(f"Condiciones de pago: {condiciones}")
        if metodo_pago:
            note_lines.append(f"Método de pago: {metodo_pago}")

        return {
            'referencia_todopintura': codigo,
            'name': name,
            'note': "<br/>".join(escape(line) for line in note_lines) if note_lines else False,
        }

    @api.model
    def _create_or_update_payment_term(self, codigo, record):
        payment_term = self._find_existing_payment_term(codigo)

        if payment_term:
            payment_term.write({key: value for key, value in record.items() if key != 'referencia_todopintura'})
            status = 'updated'
        else:
            payment_term = self.env['account.payment.term'].create(record)
            status = 'created'

        return payment_term, status

    @api.model
    def _iter_rows(self, sheet, is_xlsx):
        if is_xlsx:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                yield list(row)
        else:
            for row_index in range(1, sheet.nrows):
                yield [sheet.cell(row_index, col).value for col in range(sheet.ncols)]

    @api.model
    def _process_import_row(self, row, results):
        """Procesa una fila del Excel: crea/actualiza la forma de pago y registra el resultado."""
        def cell(idx):
            return row[idx] if len(row) > idx else None

        codigo = self._safe_int(cell(0))
        name = self._sanitize_import_text(cell(1))
        condiciones = self._sanitize_import_text(cell(2))
        metodo_pago = self._sanitize_import_text(cell(3))

        if codigo is None or not name:
            return

        row_label = f"[{codigo}] {name}"

        try:
            record = self._prepare_payment_term_record(codigo, name, condiciones, metodo_pago)
            _payment_term, status = self._create_or_update_payment_term(codigo, record)
            results[status].append(row_label)
        except Exception as exc:
            _logger.exception("Error al importar la forma de pago %s", row_label)
            results['error'].append(f"{row_label}: {exc}")

    @api.model
    def _build_import_log_html(self, results):
        def section(title, items, css_class):
            if not items:
                return f'<p><b>{escape(title)}:</b> ninguno.</p>'
            lines = ''.join(f'<li>{escape(line)}</li>' for line in items)
            return f'<p><b>{escape(title)} ({len(items)}):</b></p><ul class="{css_class}">{lines}</ul>'

        return (
            section('Formas de pago nuevas', results['created'], 'text-success')
            + section('Formas de pago actualizadas', results['updated'], 'text-info')
            + section('Errores', results['error'], 'text-danger')
        )

    def _load_sheet(self):
        ext = ''
        if self.file_name:
            ext = self.file_name.split('.')[-1].lower()
        data = base64.b64decode(self.file)
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'

        if ext == 'xlsx':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return wb.active, True
        elif ext == 'xls':
            if not xlrd:
                raise UserError("Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.")
            book = xlrd.open_workbook(file_contents=data)
            return book.sheet_by_index(0), False
        else:
            raise UserError("Formato de archivo no soportado. Usa .xls o .xlsx")

    def action_import_payment_terms(self):
        self.ensure_one()
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        sheet, is_xlsx = self._load_sheet()

        results = {'created': [], 'updated': [], 'error': []}
        for row in self._iter_rows(sheet, is_xlsx):
            self._process_import_row(row, results)

        self.write({
            'state': 'done',
            'created_count': len(results['created']),
            'updated_count': len(results['updated']),
            'error_count': len(results['error']),
            'log': self._build_import_log_html(results),
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.payment.terms.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_reset(self):
        self.ensure_one()
        self.write({
            'state': 'upload',
            'file': False,
            'file_name': False,
            'log': False,
            'created_count': 0,
            'updated_count': 0,
            'error_count': 0,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.payment.terms.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
