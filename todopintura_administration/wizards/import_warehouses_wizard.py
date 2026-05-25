from odoo import api, fields, models
import base64
from odoo.exceptions import UserError
import io
import traceback
import re
import logging

# Logger para registrar excepciones y depuración
_logger = logging.getLogger(__name__)
try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

class ImportWarehousesWizard(models.TransientModel):
    _name = 'import.warehouses.wizard'
    _description = 'Wizard para importar almacenes desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX', required=True)
    file_name = fields.Char('Nombre del archivo')

    @api.model
    def _sanitize_import_text(self, value):
        if value is None:
            return ''
        if isinstance(value, int):
            text = str(value)
        elif isinstance(value, float):
            if value.is_integer():
                text = str(int(value))
            else:
                text = str(value)
        else:
            text = value.strip() if isinstance(value, str) else str(value).strip()
        return '' if text and all(char == '*' for char in text) else text

    def _parse_excel_file(self, file_content, file_name):
        ext = ''
        if file_name:
            ext = file_name.split('.')[-1].lower()
        data = base64.b64decode(file_content)
        # Detección por cabecera si la extensión no es fiable
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'

        rows = []
        if ext == 'xlsx':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not any(val is not None for val in row):
                    continue
                rows.append([self._sanitize_import_text(val) for val in row])
        elif ext == 'xls':
            if not xlrd:
                raise UserError("Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.")
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
            for row_idx in range(1, sheet.nrows):
                row_vals = []
                has_data = False
                for col_idx in range(sheet.ncols):
                    val = sheet.cell(row_idx, col_idx).value
                    if val is not None and val != '':
                        has_data = True
                    row_vals.append(self._sanitize_import_text(val))
                if not has_data:
                    continue
                rows.append(row_vals)
        else:
            raise UserError("Formato de archivo no soportado. Usa .xls o .xlsx")

        parsed_data = []
        for idx, row in enumerate(rows):
            id_tp = row[0] if len(row) > 0 else ''
            name = row[1] if len(row) > 1 else ''

            # Parse mixed first column (e.g. "1 ALMACEN CENTRAL")
            if not name and id_tp and isinstance(id_tp, str) and ' ' in id_tp:
                parts = id_tp.strip().split(None, 1)
                if len(parts) > 1:
                    possible_id, possible_name = parts[0].strip(), parts[1].strip()
                    if possible_id.isdigit() or len(possible_id) <= 6:
                        id_tp = possible_id
                        name = possible_name

            # Search in subsequent columns if name is still empty
            if not name:
                try:
                    for extra in row[2:]:
                        if extra and not extra.isdigit() and len(extra) > 1:
                            name = extra
                            break
                except Exception:
                    pass

            if not id_tp and not name:
                continue

            parsed_data.append({
                'row_num': idx + 2,
                'id_todopinturas': id_tp,
                'name': name or (id_tp and str(id_tp).strip()) or f'Almacén {idx + 1}',
            })
        return parsed_data

    def _ensure_warehouse_locations(self, name, company_id):
        """Ensures that the view location and internal stock location exist for the warehouse name."""
        stock_loc_model = self.env['stock.location'].sudo()
        
        # Search/Create view location
        view_loc = stock_loc_model.with_context(lang=None).search([
            ('name', '=', name),
            ('usage', '=', 'view'),
            ('company_id', '=', company_id)
        ], limit=1)
        if not view_loc:
            view_loc = stock_loc_model.create({
                'name': name or 'Default View',
                'usage': 'view',
                'company_id': company_id,
            })
            
        # Search/Create internal stock location
        stock_loc = stock_loc_model.with_context(lang=None).search([
            ('name', 'ilike', name),
            ('usage', '=', 'internal'),
            ('company_id', '=', company_id)
        ], limit=1)
        if not stock_loc:
            stock_loc = stock_loc_model.create({
                'name': f"{name} Stock" if name else 'Stock',
                'usage': 'internal',
                'location_id': view_loc.id,
                'company_id': company_id,
            })
            
        return view_loc.id, stock_loc.id

    def _generate_warehouse_code(self, name, id_tp, company_id):
        base_code = (id_tp and str(id_tp).strip()) or re.sub(r"\W+", "", name or '')[:12]
        if not base_code:
            base_code = 'WH'
        code = base_code
        suffix = 0
        wh_model = self.env['stock.warehouse'].sudo()
        while wh_model.search([('code', '=', code), ('company_id', '=', company_id)], limit=1):
            suffix += 1
            code = f"{base_code}_{suffix}"
        return code

    def _create_or_update_warehouse(self, id_tp, name, company_id):
        warehouse = None
        wh_model = self.env['stock.warehouse'].sudo()
        
        # 1. Search existing warehouse
        if id_tp:
            warehouse = wh_model.with_context(active_test=False).search([
                ('id_todopinturas', '=', id_tp),
                ('company_id', '=', company_id)
            ], limit=1)
        if not warehouse and name:
            warehouse = wh_model.with_context(active_test=False).search([
                ('name', '=', name),
                ('company_id', '=', company_id)
            ], limit=1)
            
        vals = {
            'name': name,
            'company_id': company_id,
        }
        if id_tp:
            vals['id_todopinturas'] = id_tp
            
        if warehouse:
            # Update existing
            warehouse.write(vals)
            is_new = False
        else:
            # Create new
            view_loc_id, stock_loc_id = self._ensure_warehouse_locations(name, company_id)
            vals['view_location_id'] = view_loc_id
            vals['lot_stock_id'] = stock_loc_id
            
            code = self._generate_warehouse_code(name, id_tp, company_id)
            vals['code'] = code
            
            warehouse = wh_model.create(vals)
            is_new = True
            
        return warehouse, is_new

    def _configure_replenishment_route(self, warehouse, central_warehouse, company_id):
        """Creates or updates a custom direct supply route and rule from central to warehouse."""
        route_model = self.env['stock.route'].sudo()
        rule_model = self.env['stock.rule'].sudo()
        
        route_name = f"Pedir a Central -> {warehouse.name}"
        
        # 1. Find or create the route
        route = route_model.search([
            ('name', '=', route_name),
            ('company_id', '=', company_id)
        ], limit=1)
        
        route_vals = {
            'name': route_name,
            'company_id': company_id,
            'warehouse_selectable': True,
            'warehouse_ids': [(4, warehouse.id)],
        }
        
        if route:
            route.write(route_vals)
            route_created = False
        else:
            route = route_model.create(route_vals)
            route_created = True
            
        # 2. Find or create/update the rule
        src_loc_id = central_warehouse.lot_stock_id.id
        dest_loc_id = warehouse.lot_stock_id.id
        
        if not src_loc_id or not dest_loc_id:
            raise UserError(f"No se pudieron determinar las ubicaciones de stock de la central ({central_warehouse.name}) o del almacén ({warehouse.name}).")
            
        rule = rule_model.search([
            ('route_id', '=', route.id),
            ('warehouse_id', '=', warehouse.id),
            ('company_id', '=', company_id)
        ], limit=1)
        
        rule_vals = {
            'name': f'From Central to {warehouse.name}',
            'route_id': route.id,
            'location_src_id': src_loc_id,
            'location_dest_id': dest_loc_id,
            'action': 'pull',
            'procure_method': 'make_to_stock',
            'warehouse_id': warehouse.id,
            'picking_type_id': warehouse.int_type_id.id,
            'company_id': company_id,
        }
        
        if rule:
            rule.write(rule_vals)
        else:
            rule_model.create(rule_vals)
            
        return route_created

    def action_import_warehouses(self):
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")
            
        if 'stock.route' not in self.env or 'stock.rule' not in self.env:
            raise UserError("El módulo de Inventario (stock) no está instalado o no se puede acceder a él.")

        # Company actual
        company = self.env.company

        # Parse file
        parsed_data = self._parse_excel_file(self.file, self.file_name)
        if not parsed_data:
            raise UserError("El archivo no contiene almacenes válidos para importar.")

        created = 0
        updated = 0
        created_routes = 0
        updated_routes = 0
        failed_routes = []
        per_warehouse_logs = []

        central_wh = None
        for idx, wh_data in enumerate(parsed_data):
            id_tp = wh_data['id_todopinturas']
            name = wh_data['name']
            row_num = wh_data['row_num']
            
            try:
                # 1. Create or update warehouse
                warehouse, is_new = self._create_or_update_warehouse(id_tp, name, company.id)
                if is_new:
                    created += 1
                else:
                    updated += 1
                    
                # 2. Configure warehouse roles and routes
                if idx == 0:
                    # Central warehouse configuration
                    central_wh = warehouse
                    
                    # Uncheck tp_is_central_request_hub on other warehouses in this company
                    other_hubs = self.env['stock.warehouse'].sudo().search([
                        ('company_id', '=', company.id),
                        ('tp_is_central_request_hub', '=', True),
                        ('id', '!=', warehouse.id)
                    ])
                    if other_hubs:
                        other_hubs.write({'tp_is_central_request_hub': False})
                        
                    warehouse.write({
                        'tp_is_central_request_hub': True,
                        'buy_to_resupply': True,
                        'resupply_wh_ids': [(5, 0, 0)],
                    })
                    per_warehouse_logs.append(f"{id_tp or ''} - {name}: Configured as CENTRAL hub")
                else:
                    # Branch warehouse configuration
                    warehouse.write({
                        'tp_is_central_request_hub': False,
                        'buy_to_resupply': False,
                        'resupply_wh_ids': [(5, 0, 0)],
                    })
                    
                    # Create replenishment route if central is known
                    if central_wh:
                        try:
                            route_created = self._configure_replenishment_route(warehouse, central_wh, company.id)
                            if route_created:
                                created_routes += 1
                            else:
                                updated_routes += 1
                            per_warehouse_logs.append(f"{id_tp or ''} - {name}: Replenished from Central ({central_wh.name})")
                        except Exception as route_err:
                            err_msg = str(route_err)
                            failed_routes.append(f"{name}: {err_msg}")
                            _logger.exception("Error creating route/rule for warehouse %s", name)
                            per_warehouse_logs.append(f"{id_tp or ''} - {name}: ERROR creating route/rule: {err_msg}")
                    else:
                        per_warehouse_logs.append(f"{id_tp or ''} - {name}: Skip route creation (No central warehouse identified)")
                        
            except Exception as e:
                tb = traceback.format_exc()
                raise UserError(f"Error procesando la fila {row_num}: {e}\n\nTraceback:\n{tb}")

        # Summary message
        dedup_failed = list(dict.fromkeys(failed_routes))
        details = "; ".join(per_warehouse_logs) if per_warehouse_logs else ""
        message = f'Almacenes creados: {created}, actualizados: {updated}. Rutas creadas: {created_routes}, actualizadas: {updated_routes}. Errores rutas: {len(dedup_failed)}'
        if dedup_failed:
            message += f" - {'; '.join(dedup_failed)}"
        if details:
            message += f". Detalles: {details}"

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Importación finalizada',
                'message': message,
                'sticky': False,
            },
        }

