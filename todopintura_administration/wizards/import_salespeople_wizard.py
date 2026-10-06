import base64
import io
import re
import unicodedata

from markupsafe import escape

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import xlrd
except ImportError:
    xlrd = None


class ImportSalespeopleWizard(models.TransientModel):
    _name = 'import.salespeople.wizard'
    _description = 'Wizard para importar comerciales (usuarios) desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX')
    file_name = fields.Char('Nombre del archivo')
    state = fields.Selection([
        ('upload', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='upload', required=True)
    log = fields.Html('Registro de importación', readonly=True)
    created_count = fields.Integer('Comerciales nuevos', readonly=True)
    updated_count = fields.Integer('Comerciales actualizados', readonly=True)
    error_count = fields.Integer('Errores', readonly=True)

    @api.model
    def _normalize_code(self, value):
        """Código de comercial como texto ('21', no '21.0'); '' si vacío o 0."""
        if isinstance(value, (int, float)):
            return str(int(value)) if value else ''
        value = (value or '').strip()
        try:
            return str(int(float(value))) if float(value) else ''
        except ValueError:
            return value

    @api.model
    def _build_login(self, name):
        """'comercial_<nombre>' en minúsculas, sin tildes y con '_' en lugar de espacios."""
        text = unicodedata.normalize('NFKD', name.strip().lower())
        text = ''.join(char for char in text if not unicodedata.combining(char))
        return 'comercial_' + re.sub(r'\W+', '_', text).strip('_')

    def _read_rows(self):
        data = base64.b64decode(self.file)
        if data[:2] == b'PK':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx.")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return [list(row) for row in wb.active.iter_rows(min_row=2, values_only=True)]
        if not xlrd:
            raise UserError("Falta la librería xlrd para procesar archivos .xls.")
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return [[sheet.cell(r, c).value for c in range(sheet.ncols)] for r in range(1, sheet.nrows)]

    def action_import_salespeople(self):
        self.ensure_one()
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        users = self.env['res.users'].with_context(active_test=False, no_reset_password=True)
        results = {'created': [], 'updated': [], 'error': []}
        for row in self._read_rows():
            code = self._normalize_code(row[0] if row else None)
            name = (str(row[1]).strip() if len(row) > 1 and row[1] else '')
            if not code or not name:
                continue
            label = f"[{code}] {name}"
            try:
                user = users.search([('commercial_code', '=', code)], limit=1)
                with self.env.cr.savepoint():
                    if user:
                        user.write({'name': name})
                        results['updated'].append(label)
                    else:
                        users.create({
                            'name': name,
                            'login': self._build_login(name),
                            'commercial_code': code,
                        })
                        results['created'].append(label)
            except Exception as exc:
                results['error'].append(f"{label}: {exc}")

        def section(title, items, css):
            if not items:
                return f'<p><b>{escape(title)}:</b> ninguno.</p>'
            lines = ''.join(f'<li>{escape(i)}</li>' for i in items)
            return f'<p><b>{escape(title)} ({len(items)}):</b></p><ul class="{css}">{lines}</ul>'

        self.write({
            'state': 'done',
            'created_count': len(results['created']),
            'updated_count': len(results['updated']),
            'error_count': len(results['error']),
            'log': (section('Comerciales nuevos', results['created'], 'text-success')
                    + section('Comerciales actualizados', results['updated'], 'text-info')
                    + section('Errores', results['error'], 'text-danger')),
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
