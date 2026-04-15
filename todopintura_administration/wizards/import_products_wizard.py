from odoo import fields, models
import base64
import io
import logging

try:
    import xlrd
except ImportError:
    xlrd = None

from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
except ImportError:
    openpyxl = None


_logger = logging.getLogger(__name__)


class ImportProductsWizard(models.TransientModel):
    _name = 'import.products.wizard'
    _description = 'Wizard para importar productos desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX', required=True)
    file_name = fields.Char('Nombre del archivo')
    error_log = fields.Text('Errores', readonly=True)

    def _safe_int(self, value, default=None):
        if value in (None, False, ''):
            return default
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return default
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default

    def _safe_float(self, value, default=None):
        if value in (None, False, ''):
            return default
        if isinstance(value, str):
            value = value.strip().replace(',', '.')
            if not value:
                return default
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    def _get_cell(self, row, idx):
        return row[idx] if idx < len(row) else None

    def _ensure_default_tariffs(self):
        tariffs = {}
        for index in range(1, 8):
            name = f"Tarifa {index}"
            tariff = self.env['product.pricelist'].search([('name', '=', name)], limit=1)
            if not tariff:
                tariff = self.env['product.pricelist'].create({'name': name, 'company_id': False})
                _logger.info("Creada tarifa base %s", name)
            tariffs[name] = tariff
        return tariffs

    def _iter_rows(self, sheet, is_xlsx):
        if is_xlsx:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                yield tuple(row)
            return
        for row_idx in range(1, sheet.nrows):
            yield tuple(sheet.row_values(row_idx))

    def _get_barcode_domain(self, barcode, company_id=False):
        domain = [('barcode', '=', barcode)]
        if company_id:
            domain.append(('company_id', 'in', (False, company_id)))
        return domain

    def _has_barcode_conflict(self, barcode, product=False):
        if not barcode:
            return False

        company_id = product.company_id.id if product and product.company_id else False
        product_domain = self._get_barcode_domain(barcode, company_id=company_id)
        packaging_domain = self._get_barcode_domain(barcode, company_id=company_id)

        if product and product.product_variant_ids:
            product_domain.append(('id', 'not in', product.product_variant_ids.ids))

        if self.env['product.product'].sudo().search_count(product_domain, limit=1):
            return True

        return bool(self.env['product.uom'].sudo().search_count(packaging_domain, limit=1))

    def _sanitize_barcode(self, record, product=False):
        barcode = record.get('barcode')
        if barcode and self._has_barcode_conflict(barcode, product=product):
            _logger.warning(
                "Código de barras duplicado %s para producto %s. Se deja vacío para evitar el error de validación.",
                barcode,
                record.get('default_code'),
            )
            print(
                f"Código de barras duplicado {barcode} para producto {record.get('default_code')}. "
                "Se deja vacío para evitar el error de validación."
            )
            record['barcode'] = None
        return record

    def _prepare_product_values(self, row):
        get_cell = self._get_cell

        num_prod = self._safe_int(get_cell(row, 0), default=0)
        try:
            name = str(get_cell(row, 1)).strip() if get_cell(row, 1) is not None else ''
        except ValueError:
            name = ''
        taxes_id_name = self._safe_int(get_cell(row, 2), default=0)
        barcode_value = get_cell(row, 23)
        if isinstance(barcode_value, float):
            barcode = str(int(barcode_value)).strip()
        else:
            barcode = str(barcode_value).strip() if barcode_value else ''
        description = str(get_cell(row, 39)).strip() if get_cell(row, 39) is not None else ''
        coste = self._safe_float(get_cell(row, 24), default=0.0) or 0.0
        art_prov = str(get_cell(row, 22)).strip() if get_cell(row, 22) is not None else ''
        pos_categ_ref = self._safe_int(get_cell(row, 28), default=None)
        num_prov = self._safe_int(get_cell(row, 30), default=None)
        price_last_buy = self._safe_float(get_cell(row, 24), default=None)
        invoice_description = str(get_cell(row, 20)).strip() if get_cell(row, 20) is not None else ''

        observations = [
            description,
            art_prov,
            str(get_cell(row, 40)).strip() if get_cell(row, 40) is not None else '',
            str(get_cell(row, 41)).strip() if get_cell(row, 41) is not None else '',
            str(get_cell(row, 42)).strip() if get_cell(row, 42) is not None else '',
            str(get_cell(row, 43)).strip() if get_cell(row, 43) is not None else '',
            str(get_cell(row, 44)).strip() if get_cell(row, 44) is not None else '',
            str(get_cell(row, 45)).strip() if get_cell(row, 45) is not None else '',
            str(get_cell(row, 46)).strip() if get_cell(row, 46) is not None else '',
            str(get_cell(row, 47)).strip() if get_cell(row, 47) is not None else '',
            str(get_cell(row, 48)).strip() if get_cell(row, 48) is not None else '',
        ]
        notes = "<br/>".join(filter(None, observations))

        prices = []
        discounts = []
        for idx in range(3, 17, 2):
            prices.append(self._safe_float(get_cell(row, idx), default=None))
            discounts.append(self._safe_float(get_cell(row, idx + 1), default=None))

        return {
            'num_prod': num_prod,
            'name': name,
            'taxes_id_name': taxes_id_name,
            'barcode': barcode,
            'description': description,
            'coste': coste,
            'pos_categ_ref': pos_categ_ref,
            'num_prov': num_prov,
            'price_last_buy': price_last_buy,
            'invoice_description': invoice_description,
            'notes': notes,
            'prices': prices,
            'discounts': discounts,
        }

    def action_import_products(self):
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        data = base64.b64decode(self.file)
        ext = ''
        if self.file_name:
            ext = self.file_name.split('.')[-1].lower()
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'
        print(f"Extensión detectada: {ext}")
        sheet = None
        is_xlsx = False
        if ext == 'xlsx':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.")
            is_xlsx = True
            print("Usando openpyxl para .xlsx")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = wb.active
        elif ext == 'xls':
            if not xlrd:
                raise UserError("Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.")
            print("Usando xlrd para .xls")
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
        else:
            raise UserError("Formato de archivo no soportado. Usa .xls o .xlsx")

        errors = []
        tariff_names = [f"Tarifa {i}" for i in range(1, 8)]
        tariffs = self._ensure_default_tariffs()
        print(f"Procesando filas con {'openpyxl' if is_xlsx else 'xlrd'}...")

        for row_idx, row in enumerate(self._iter_rows(sheet, is_xlsx)):
            excel_row = row_idx + 2
            print(f"Fila {excel_row} ({'xlsx' if is_xlsx else 'xls'}): {row}")
            values = self._prepare_product_values(row)

            if not (
                values['num_prod']
                or values['name']
                or values['taxes_id_name']
                or values['barcode']
                or values['description']
            ):
                print("Todos los datos están vacíos. Terminando la importación.")
                break

            pos_categ = self.env['pos.category'].search(
                [('referencia_todopintura', '=', values['pos_categ_ref'])], limit=1
            ) if values['pos_categ_ref'] else False
            print(f"pos_categ: {pos_categ}")

            partner = None
            if values['num_prov']:
                print(f"Buscando proveedor con referencia: 0{values['num_prov']}")
                partner = self.env['res.partner'].search([('ref', '=', f"0{values['num_prov']}")], limit=1)
                if partner:
                    print(f"Proveedor encontrado: {partner.name} (ID: {partner.id})")
                else:
                    print(f"No se encontró proveedor con referencia 0{values['num_prov']}")

            tax_id = self.env['account.tax'].search([('name', '=', '21% G')], limit=1).id
            list_price = next((price for price in values['prices'] if price is not None), 0.0)
            record = {
                'default_code': values['num_prod'],
                'name': values['name'],
                'list_price': list_price,
                'taxes_id': [(6, 0, [tax_id])] if tax_id else [],
                'barcode': values['barcode'] if values['barcode'] else None,
                'description': values['notes'] if values['notes'] else None,
                'standard_price': values['coste'] if values['coste'] else 0,
                'invoice_description': values['invoice_description'] if values['invoice_description'] else None,
                'type': "consu",
                'is_storable': True,
                'invoice_policy': "delivery",
                'available_in_pos': True,
                'pos_categ_ids': [(6, 0, [pos_categ.id])] if pos_categ else [],
            }
            print("*" * 40)
            print(f"Producto: {values['name']} (ID: {values['num_prod']})")
            print("*" * 40)

            supplier_partner = partner if (values['num_prov'] and values['price_last_buy']) else None
            if pos_categ:
                category = self.crear_categoria_con_padres(pos_categ, supplier_partner)
                record['categ_id'] = category.id
                print(f"category: {category}")
            elif supplier_partner:
                category = self.env['product.category'].search([('name', '=', supplier_partner.name)], limit=1)
                if not category:
                    category = self.env['product.category'].create({'name': supplier_partner.name})
                record['categ_id'] = category.id

            product = self._create_or_update_product(record)
            if not product:
                errors.append(f"Fila {excel_row}: no se pudo crear o actualizar el producto {values['name']!r}.")
                continue

            product_variant = product.product_variant_id
            if product_variant and values['num_prov'] and values['price_last_buy'] and partner:
                supplierinfo = self.env['product.supplierinfo'].search([
                    ('product_id', '=', product_variant.id),
                    ('partner_id', '=', partner.id)
                ], limit=1)
                supplierinfo_vals = {
                    'partner_id': partner.id,
                    'product_id': product_variant.id,
                    'price': values['price_last_buy'],
                }
                if supplierinfo:
                    supplierinfo.write(supplierinfo_vals)
                    print("Proveedor actualizado")
                else:
                    self.env['product.supplierinfo'].create(supplierinfo_vals)
                    print("Proveedor creado")

            for idx, tariff_name in enumerate(tariff_names):
                self.create_or_update_tariffs(
                    product,
                    values['prices'][idx] if idx < len(values['prices']) else None,
                    values['discounts'][idx] if idx < len(values['discounts']) else None,
                    tariff_name,
                    tariffs=tariffs,
                )

        if errors:
            self.error_log = "\n".join(errors)

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.products.wizard',
            'view_mode': 'form',
            'res_id': self.id,
        }

    def _create_or_update_product(self, record):
        if record['name']:
            product = self.env['product.template'].search([('default_code', '=', record['default_code'])], limit=1)

            if product:
                record = self._sanitize_barcode(record, product=product)
                try:
                    product.write(record)
                except ValidationError as exc:
                    if record.get('barcode') and 'barcode' in str(exc).lower():
                        _logger.warning(
                            "Error de barcode duplicado al actualizar %s. Reintentando con barcode vacío.",
                            record.get('default_code'),
                        )
                        record['barcode'] = None
                        product.write(record)
                    else:
                        raise
            else:
                record = self._sanitize_barcode(record)
                try:
                    product = self.env['product.template'].create(record)
                except ValidationError as exc:
                    if record.get('barcode') and 'barcode' in str(exc).lower():
                        _logger.warning(
                            "Error de barcode duplicado al crear %s. Reintentando con barcode vacío.",
                            record.get('default_code'),
                        )
                        record['barcode'] = None
                        product = self.env['product.template'].create(record)
                    else:
                        raise
            return product

    def create_or_update_tariffs(self, product, price, descuento, tariff_name, tariffs=None):
        if not product:
            return

        tariff = tariffs.get(tariff_name) if tariffs else self.env['product.pricelist'].search([('name', '=', tariff_name)], limit=1)
        if not tariff:
            tariff = self.env['product.pricelist'].create({'name': tariff_name, 'company_id': False})

        pricelist_items = self.env['product.pricelist.item'].search([
            ('pricelist_id', '=', tariff.id),
            ('product_tmpl_id', '=', product.id)
        ])

        if descuento in (None, False, ''):
            _logger.info(
                "No se crea línea para %s del producto %s porque no tiene descuento válido: %s",
                tariff_name,
                product.default_code,
                descuento,
            )
            if pricelist_items:
                pricelist_items.unlink()
            return

        try:
            descuento = float(descuento)
        except (TypeError, ValueError):
            _logger.info(
                "No se crea línea para %s del producto %s porque el descuento no es numérico: %s",
                tariff_name,
                product.default_code,
                descuento,
            )
            if pricelist_items:
                pricelist_items.unlink()
            return

        if descuento == 0:
            _logger.info(
                "No se crea línea para %s del producto %s porque no tiene descuento aplicable: %s",
                tariff_name,
                product.default_code,
                descuento,
            )
            if pricelist_items:
                pricelist_items.unlink()
            return

        pricelist_item_vals = {
            'pricelist_id': tariff.id,
            'product_tmpl_id': product.id,
            'applied_on': '1_product',
            'compute_price': 'percentage',
            'percent_price': descuento,
        }

        pricelist_item = pricelist_items[:1]
        if pricelist_item:
            (pricelist_items - pricelist_item).unlink()
            print("Producto actualizado")
            pricelist_item.write(pricelist_item_vals)
        else:
            self.env['product.pricelist.item'].create(pricelist_item_vals)

    def crear_categoria_con_padres(self, pos_categ, partner=None):
        """
        Crea una estructura de categorías donde:
        1. Si hay proveedor, crea/busca una categoría con su nombre
        2. Crea una subcategoría específica para este proveedor+categoría
        """
        if not partner:
            category = self.env['product.category'].search([('name', '=', pos_categ.name)], limit=1)
            if not category:
                category = self.env['product.category'].create({'name': pos_categ.name})
                print(f"Categoría creada sin proveedor: {pos_categ.name}")
            else:
                print(f"Usando categoría existente: {pos_categ.name}")
            parent_pos = pos_categ.parent_id
            if parent_pos and not category.parent_id:
                parent_category = self.crear_categoria_con_padres(parent_pos)
                category.write({'parent_id': parent_category.id})
                print(f"Actualizada jerarquía para: {category.name}")
            return category
        provider_category = self.env['product.category'].search([('name', '=', partner.name)], limit=1)
        if not provider_category:
            provider_category = self.env['product.category'].create({'name': partner.name})
            print(f"Creada categoría de proveedor: {partner.name}")
        specific_category = self.env['product.category'].search([
            ('name', '=', pos_categ.name),
            ('parent_id', '=', provider_category.id)
        ], limit=1)
        if not specific_category:
            specific_category = self.env['product.category'].create({
                'name': pos_categ.name,
                'parent_id': provider_category.id
            })
            print(f"Creada subcategoría: {pos_categ.name} bajo proveedor: {partner.name}")
        else:
            print(f"Usando categoría existente: {pos_categ.name} bajo proveedor: {partner.name}")
        return specific_category
