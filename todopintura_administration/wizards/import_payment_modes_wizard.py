import base64
import io
import logging
import re

from markupsafe import escape

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

_logger = logging.getLogger(__name__)


class ImportPaymentModesWizard(models.TransientModel):
    _name = 'import.payment.modes.wizard'
    _description = 'Wizard para importar modos de pago desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX')
    file_name = fields.Char('Nombre del archivo')
    state = fields.Selection([
        ('upload', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='upload', required=True)
    log = fields.Html('Registro de importación', readonly=True)
    created_count = fields.Integer('Modos de pago nuevos', readonly=True)
    updated_count = fields.Integer('Modos de pago actualizados', readonly=True)
    error_count = fields.Integer('Errores', readonly=True)

    _iban_re = re.compile(r'(ES\d{22})', re.IGNORECASE)

    @api.model
    def _sanitize_import_text(self, value):
        if value is None:
            return ''
        text = value.strip() if isinstance(value, str) else str(value).strip()
        return '' if text and all(char == '*' for char in text) else text

    @api.model
    def _iter_rows(self, sheet, is_xlsx):
        if is_xlsx:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                yield list(row)
        else:
            for row_index in range(1, sheet.nrows):
                yield [sheet.cell(row_index, col).value for col in range(sheet.ncols)]

    @api.model
    def _find_existing_payment_mode(self, name):
        return self.env['account.payment.mode'].with_context(active_test=False).search([
            ('name', '=', name), ('company_id', '=', self.env.company.id)], limit=1)

    @api.model
    def _get_or_create_payment_method(self, name):
        """Find or create an account.payment.method to link from account.payment.mode."""
        ppm = self.env['account.payment.method'].with_context(active_test=False).search([('name', '=', name)], limit=1)
        if ppm:
            return ppm
        # create a code from name: uppercase letters and digits
        code = ''.join(ch for ch in name.upper() if ch.isalnum())[:12] or 'AUTO'
        # ensure unique code by appending a number if needed
        base_code = code
        i = 1
        while self.env['account.payment.method'].search([('code', '=', code), ('payment_type', '=', 'inbound')]):
            code = f"{base_code[:10]}{i}"
            i += 1
        vals = {'name': name, 'code': code, 'payment_type': 'inbound'}
        return self.env['account.payment.method'].create(vals)

    @api.model
    def _create_or_update_payment_mode(self, name, iban):
        payment_mode = self._find_existing_payment_mode(name)
        # ensure there is an account.payment.method to link
        payment_method = self._get_or_create_payment_method(name)
        vals = {
            'name': name,
            'company_id': self.env.company.id,
            'payment_method_id': payment_method.id,
            'bank_account_link': 'fixed' if iban else 'variable',
        }
        if payment_mode:
            # update fields if necessary
            update_vals = {}
            if payment_mode.payment_method_id.id != payment_method.id:
                update_vals['payment_method_id'] = payment_method.id
            if payment_mode.bank_account_link != vals['bank_account_link']:
                update_vals['bank_account_link'] = vals['bank_account_link']
            if update_vals:
                payment_mode.write(update_vals)
            status = 'updated'
        else:
            payment_mode = self.env['account.payment.mode'].create(vals)
            status = 'created'
        return payment_mode, status

    @api.model
    def _process_import_row(self, row, results):
        def cell(idx):
            return row[idx] if len(row) > idx else None

        metodo_pago_raw = self._sanitize_import_text(cell(3))
        if not metodo_pago_raw:
            return

        # Normalize and extract IBAN
        text = metodo_pago_raw.strip()
        match = self._iban_re.search(text)
        iban = match.group(1).upper() if match else False
        if match:
            # remove iban from text
            name = (text[:match.start()] + text[match.end():]).strip()
        else:
            name = text

        name = name.upper()
        row_label = name if name else metodo_pago_raw

        try:
            payment_mode, status = self._create_or_update_payment_mode(name, iban)
            results[status].append(row_label)
        except Exception as exc:
            _logger.exception('Error al importar el modo de pago %s', row_label)
            results['error'].append(f"{row_label}: {exc}")

    @api.model
    def _build_import_log_html(self, results):
        def section(title, items, css_class):
            if not items:
                return f'<p><b>{escape(title)}:</b> ninguno.</p>'
            lines = ''.join(f'<li>{escape(line)}</li>' for line in items)
            return f'<p><b>{escape(title)} ({len(items)}):</b></p><ul class="{css_class}">{lines}</ul>'

        return (
            section('Modos de pago nuevos', results['created'], 'text-success')
            + section('Modos de pago actualizados', results['updated'], 'text-info')
            + section('Errores', results['error'], 'text-danger')
        )

    def _load_sheet(self):
        ext = ''
        if self.file_name:
            ext = self.file_name.split('.')[-1].lower()
        data = base64.b64decode(self.file)
        if not ext or ext not in ['xls', 'xlsx']:
            if data[:2] == b'PK':
                ext = 'xlsx'
            elif data[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                ext = 'xls'

        if ext == 'xlsx':
            if not openpyxl:
                raise UserError('Falta la librería openpyxl para procesar archivos .xlsx. Por favor, instálala.')
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return wb.active, True
        elif ext == 'xls':
            if not xlrd:
                raise UserError('Falta la librería xlrd para procesar archivos .xls. Por favor, instálala.')
            book = xlrd.open_workbook(file_contents=data)
            return book.sheet_by_index(0), False
        else:
            raise UserError('Formato de archivo no soportado. Usa .xls o .xlsx')

    def action_import_payment_modes(self):
        self.ensure_one()
        if not self.file:
            raise UserError('Por favor, sube un archivo XLS o XLSX.')

        sheet, is_xlsx = self._load_sheet()

        results = {'created': [], 'updated': [], 'error': []}
        for row in self._iter_rows(sheet, is_xlsx):
            self._process_import_row(row, results)

        self.write({
            'state': 'done',
            'created_count': len(results['created']),
            'updated_count': len(results['updated']),
            'error_count': len(results['error']),
            'log': self._build_import_log_html(results),
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.payment.modes.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_reset(self):
        self.ensure_one()
        self.write({
            'state': 'upload',
            'file': False,
            'file_name': False,
            'log': False,
            'created_count': 0,
            'updated_count': 0,
            'error_count': 0,
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.payment.modes.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }



