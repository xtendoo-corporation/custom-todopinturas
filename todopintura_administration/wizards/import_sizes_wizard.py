import base64
import io

from markupsafe import escape

from odoo import fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import xlrd
except ImportError:
    xlrd = None


class ImportSizesWizard(models.TransientModel):
    _name = 'import.sizes.wizard'
    _description = 'Wizard para importar tamaños de producto desde un archivo XLS o XLSX'

    file = fields.Binary('Subir archivo XLS o XLSX')
    file_name = fields.Char('Nombre del archivo')
    state = fields.Selection([
        ('upload', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='upload', required=True)
    log = fields.Html('Registro de importación', readonly=True)
    created_count = fields.Integer('Tamaños nuevos', readonly=True)
    updated_count = fields.Integer('Tamaños actualizados', readonly=True)
    error_count = fields.Integer('Errores', readonly=True)

    def _read_rows(self):
        data = base64.b64decode(self.file)
        if data[:2] == b'PK':
            if not openpyxl:
                raise UserError("Falta la librería openpyxl para procesar archivos .xlsx.")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return [list(row) for row in wb.active.iter_rows(min_row=2, values_only=True)]
        if not xlrd:
            raise UserError("Falta la librería xlrd para procesar archivos .xls.")
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return [[sheet.cell(r, c).value for c in range(sheet.ncols)] for r in range(1, sheet.nrows)]

    def action_import_sizes(self):
        self.ensure_one()
        if not self.file:
            raise UserError("Por favor, sube un archivo XLS o XLSX.")

        sizes_by_code = {size.code: size for size in self.env['product.size'].search([])}
        results = {'created': [], 'updated': [], 'error': []}
        for row in self._read_rows():
            try:
                code = int(float(row[0])) if row and row[0] not in (None, '') else None
            except (TypeError, ValueError):
                code = None
            name = ' '.join(str(row[1]).split()) if len(row) > 1 and row[1] else ''
            # Código 0 ("--") significa "sin tamaño": no se importa.
            if not code or not name:
                continue
            label = f"[{code}] {name}"
            try:
                size = sizes_by_code.get(code)
                if size:
                    size.name = name
                    results['updated'].append(label)
                else:
                    sizes_by_code[code] = self.env['product.size'].create({'code': code, 'name': name})
                    results['created'].append(label)
            except Exception as exc:
                results['error'].append(f"{label}: {exc}")

        def section(title, items, css):
            if not items:
                return f'<p><b>{escape(title)}:</b> ninguno.</p>'
            lines = ''.join(f'<li>{escape(i)}</li>' for i in items)
            return f'<p><b>{escape(title)} ({len(items)}):</b></p><ul class="{css}">{lines}</ul>'

        self.write({
            'state': 'done',
            'created_count': len(results['created']),
            'updated_count': len(results['updated']),
            'error_count': len(results['error']),
            'log': (section('Tamaños nuevos', results['created'], 'text-success')
                    + section('Tamaños actualizados', results['updated'], 'text-info')
                    + section('Errores', results['error'], 'text-danger')),
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
